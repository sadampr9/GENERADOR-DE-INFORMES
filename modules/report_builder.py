import os
import shutil
import logging
import re
import glob
from datetime import datetime
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from config import CHART_DESCRIPTIONS

logger = logging.getLogger("reniec_anomalias")

REVISION_DIR = r"C:\Users\sperezr\Documents\Proyectos - Proyectos\REVISION SEMANAL DE CONSUMOS"

def set_cell_background(cell, color_hex):
    """
    Establece el color de fondo de una celda de tabla en Word.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tc_pr.append(shd)

def replace_placeholder(doc, placeholder, replacement):
    """
    Reemplaza de forma segura un placeholder en todos los párrafos y tablas del documento.
    """
    for paragraph in doc.paragraphs:
        if placeholder in paragraph.text:
            for run in paragraph.runs:
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, replacement)
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, replacement)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if placeholder in paragraph.text:
                        for run in paragraph.runs:
                            if placeholder in run.text:
                                run.text = run.text.replace(placeholder, replacement)
                        if placeholder in paragraph.text:
                            paragraph.text = paragraph.text.replace(placeholder, replacement)

def find_historical_chart(empresa_name, ruc, servicio):
    """
    Busca el gráfico histórico 2025+ en el proyecto REVISION SEMANAL DE CONSUMOS.
    """
    graficos_dir = os.path.join(REVISION_DIR, "output", "graficos")
    if not os.path.exists(graficos_dir):
        return None
        
    all_files = os.listdir(graficos_dir)
    emp_clean = re.sub(r'[^a-zA-Z0-9]', '', str(empresa_name).lower())
    
    alias_map = {
        "redperuanadepagossac": ["red_peruana_pagos", "red_peruana"],
        "redperuanadepagos": ["red_peruana_pagos", "red_peruana"],
        "redperuana": ["red_peruana_pagos", "red_peruana"],
        "pagosdigitales": ["pagos_digitales", "pagosdigitales"],
        "bancodelanacion": ["nacion"],
        "bancodecredito": ["bcp"],
        "afpintegra": ["afp_integra"],
        "bancofalabella": ["falabella"],
        "conareme": ["conareme"],
        "xploring": ["xploring"],
        "sanmartin": ["coop_san_martin"],
        "abaco": ["abaco", "coop_abaco"],
        "clinicainternacional": ["clinica_internacional", "internacional"],
        "rimac": ["rimac"],
        "vallejo": ["cesar_vallejo", "vallejo"],
        "comercio": ["banco_comercio", "comercio"],
        "protecta": ["protecta"],
        "entel": ["entel"],
        "inei": ["inei"],
        "guinea": ["guinea"],
        "osce": ["osce"],
        "ministeriodeeducacion": ["minedu", "ministerio_educacion", "ministerio_de_educacion"],
        "minedu": ["minedu", "ministerio_educacion", "ministerio_de_educacion"],
        "sunarp": ["sunarp", "registros_publicos"],
        "registrospublicos": ["sunarp", "registros_publicos"],
        # SUCAMEC: sin gráfico histórico propio en REVISION SEMANAL — usar gráfico generado
        "sucamec": [],
        "armunicionesexplosivos": [],
        "controldeserviciosdeseguridad": [],
        "osiptel": [],
        "organismosupervisordeinversionprivadaentelecomunicaciones": [],
    }
    
    # RUCs sin gráfico histórico propio — retornar None directamente
    NO_HIST_CHART_RUCS = {"20551964940", "20216072155"}  # SUCAMEC, OSIPTEL
    if str(ruc) in NO_HIST_CHART_RUCS:
        return None
    
    search_terms = []
    for key, terms in alias_map.items():
        if key in emp_clean:
            search_terms.extend(terms)
            
    if not search_terms:
        words = [w for w in re.split(r'\W+', str(empresa_name).lower()) if len(w) > 3 and w not in ["banco", "peru", "sa", "sac", "del"]]
        search_terms.extend(words)

    candidates = []
    for fname in all_files:
        if not fname.lower().endswith(".png"):
            continue
        fname_low = fname.lower()
        for term in search_terms:
            if term in fname_low:
                candidates.append(fname)
                break
                
    if not candidates:
        return None
        
    for c in candidates:
        if "desde_2025" in c.lower() or "2025_2026" in c.lower():
            return os.path.join(graficos_dir, c)
            
    for c in candidates:
        if "linea_dedicada" in c.lower() or "julio_2026" in c.lower():
            return os.path.join(graficos_dir, c)
            
    return os.path.join(graficos_dir, candidates[0])

def get_historical_metrics(empresa_name, target_ruc, target_service, current_df=None):
    """
    Obtiene las métricas históricas de consumo (VNP 2025-2026, VID, IP%, peak_date)
    de forma dinámica a partir de la serie histórica completa (SADAM) y del mes evaluado.
    """
    from modules.chart_generator import load_sadam_full_series

    sadam_s = load_sadam_full_series(empresa_name, target_ruc, target_service)
    vnp = None
    if sadam_s is not None and not sadam_s.empty:
        vnp = float(sadam_s.mean())

    if current_df is not None and not current_df.empty:
        df_copy = current_df.copy()
        df_copy["fecha_dia"] = df_copy["fecha_consulta"].dt.date
        daily_counts = df_copy.groupby("fecha_dia").size()
        if not daily_counts.empty:
            peak_date = daily_counts.idxmax()
            vid = int(daily_counts.max())
            if vnp is None:
                vnp = float(daily_counts.mean())
            ip_pct = ((vid - vnp) / vnp * 100) if vnp > 0 else 0.0
            return {"vnp": round(vnp, 1), "vid": vid, "ip_pct": round(ip_pct, 1), "peak_date": peak_date}

    if sadam_s is not None and not sadam_s.empty:
        peak_date = sadam_s.idxmax().date()
        vid = int(sadam_s.max())
        ip_pct = ((vid - vnp) / vnp * 100) if vnp > 0 else 0.0
        return {"vnp": round(vnp, 1), "vid": vid, "ip_pct": round(ip_pct, 1), "peak_date": peak_date}

    return {"vnp": 0.0, "vid": 0, "ip_pct": 0.0, "peak_date": datetime.now().date()}

def get_service_user_count(empresa_name, service_name, ruc=None, filepath="USUARIOS POR ENTIDAD.xlsx"):
    """
    Interactúa con el archivo USUARIOS POR ENTIDAD.xlsx para obtener la cantidad
    total de usuarios asignados/registrados a la entidad y servicio.
    """
    path = Path(filepath)
    if not path.exists():
        from config import DATA_DIR
        data_path = DATA_DIR / path.name
        if data_path.exists():
            filepath = str(data_path)
        else:
            return None
    try:
        import pandas as pd
        df = pd.read_excel(filepath, usecols=[0, 2, 3, 4], header=None)
        df_data = df.iloc[6:].copy()
        
        sub_df = pd.DataFrame()
        
        # 1. Intentar filtrado por RUC exacto (10 u 11 dígitos)
        if ruc:
            ruc_clean = re.sub(r'\D', '', str(ruc))
            if len(ruc_clean) >= 10:
                ruc_prefix = ruc_clean[:10]
                mask_ruc = df_data[3].astype(str).str.contains(ruc_prefix, case=False, na=False)
                ruc_sub = df_data[mask_ruc]
                if not ruc_sub.empty:
                    sub_df = ruc_sub

        # 2. Si no hubo coincidencia por RUC, buscar por Razón Social / Empresa
        if sub_df.empty and empresa_name:
            emp_clean = re.sub(r'[^a-zA-Z0-9]', '', str(empresa_name).lower())
            alias_map = {
                'ministeriodeeducacion': 'MINISTERIO DE EDUCACION',
                'minedu': 'MINISTERIO DE EDUCACION',
                'bancodelanacion': 'BANCO DE LA NACION',
                'nacion': 'BANCO DE LA NACION',
                'bancodecredito': 'BANCO DE CREDITO',
                'bcp': 'BANCO DE CREDITO',
                'pagosdigitales': 'PAGOS DIGITALES',
                'afpintegra': 'INTEGRA',
                'autofin': 'AUTOFIN'
            }
            target_term = alias_map.get(emp_clean, None)
            
            if target_term:
                mask_emp = df_data[2].astype(str).str.contains(target_term, case=False, na=False)
                sub_df = df_data[mask_emp]

            if sub_df.empty:
                stopwords = {
                    'MINISTERIO', 'SUPERINTENDENCIA', 'MUNICIPALIDAD', 'BANCO', 'FINANCIERA',
                    'GOBIERNO', 'REGIONAL', 'DISTRITAL', 'PERU', 'S.A.C.', 'S.A.', 'SAC', 'SA',
                    'DEL', 'DE', 'LA', 'LOS', 'LAS', 'Y', 'PARA', 'POR', 'EN', 'DISTRITO', 'PROVINCIA',
                    'ORGANISMO', 'SUPERVISOR', 'NACIONAL', 'CENTRAL', 'SECRETARIA'
                }
                emp_upper = str(empresa_name).upper()
                
                if 'AFP' in emp_upper and 'BANCA' in emp_upper:
                    keywords = ['BANCA', 'SEGUROS']
                else:
                    keywords = [w for w in re.split(r'\W+', emp_upper) if len(w) >= 3 and w not in stopwords]
                
                if keywords:
                    mask_kw = pd.Series(True, index=df_data.index)
                    for kw in keywords:
                        mask_kw = mask_kw & df_data[2].astype(str).str.upper().str.contains(kw, na=False)
                    
                    sub_df = df_data[mask_kw]

        if sub_df.empty:
            return None
            
        # 3. Filtrar por Servicio específico
        if service_name and not sub_df.empty:
            serv_str = str(service_name).upper()
            if 'DEDICADA' in serv_str or 'LD' in serv_str:
                serv_pattern = 'LINEA DEDICADA|LD'
            elif 'INTERNET' in serv_str:
                serv_pattern = 'INTERNET'
            elif 'BIOMETR' in serv_str:
                serv_pattern = 'BIOMETR'
            elif 'PORTAL' in serv_str:
                serv_pattern = 'PORTAL'
            else:
                serv_pattern = re.escape(serv_str)
                
            serv_mask = sub_df[4].astype(str).str.contains(serv_pattern, case=False, na=False)
            if serv_mask.any():
                sub_df = sub_df[serv_mask]
                
        return sub_df[0].nunique()
    except Exception as e:
        logger.warning(f"Error al obtener usuarios de USUARIOS POR ENTIDAD.xlsx: {e}")
        return None

def format_heading_left(p):
    """
    Elimina sangrías y asegura alineación a la izquierda al mismo nivel que 4.2.
    """
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_top_bursts_table(doc, df, report, top_n=10):
    """
    Agrega una tabla estilizada con el Top N de ráfagas inusuales en ventanas de 60s (5 columnas, sin Usuario Operador).
    """
    if report.bursts is None or report.bursts.empty:
        return

    top_bursts = report.bursts.sort_values(by="count", ascending=False).head(top_n).reset_index(drop=True)
    if top_bursts.empty:
        return

    p_table_title = doc.add_paragraph()
    format_heading_left(p_table_title)
    run_tt = p_table_title.add_run(f"Tabla: Top {len(top_bursts)} Ráfagas de Mayor Intensidad en 60 Segundos")
    run_tt.bold = True
    p_table_title.paragraph_format.space_before = Pt(8)
    p_table_title.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1 + len(top_bursts), cols=5)
    table.style = 'Table Grid'
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["N°", "Fecha y Hora (Inicio Ventana)", "DNI / Código", "Consultas en 60s", "Velocidad (cons/seg)"]
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1F497D")
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)

    col_widths = [Inches(0.5), Inches(2.4), Inches(1.5), Inches(1.3), Inches(1.4)]

    for idx, row in top_bursts.iterrows():
        row_cells = table.rows[idx + 1].cells
        cod = row["cod_usuario"]
        dt_val = row["window_start"]
        dt_str = dt_val.strftime("%Y-%m-%d %H:%M:%S") if hasattr(dt_val, "strftime") else str(dt_val)
        cnt = int(row["count"])
        c_sec = cnt / 60.0

        row_cells[0].text = str(idx + 1)
        row_cells[1].text = dt_str
        row_cells[2].text = str(cod)
        row_cells[3].text = f"{cnt:,}"
        row_cells[4].text = f"{c_sec:.2f}"

        alignments = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
        for i, cell in enumerate(row_cells):
            for paragraph in cell.paragraphs:
                paragraph.alignment = alignments[i]
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if idx % 2 == 1:
                set_cell_background(cell, "F2F2F2")

    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(12)

def obtener_hallazgo_dinamico(chart_id, df, metadata, report, peak_date, peak_count, vnp, ip_pct, peak_user_name, peak_user_cod, image_num=1):
    """
    Genera el hallazgo forense específico de forma dinámica para cada gráfico,
    terminando explícitamente en (ver imagen X).
    """
    if chart_id == "G1":
        return (
            f"En el caso particular de la entidad analizada, el volumen diario de consultas se mantuvo en niveles bajos "
            f"hasta registrar picos históricos en los últimos meses que bordean los {int(peak_count):,} consultas, teniendo el día "
            f"{(peak_date.strftime('%d/%m/%Y') if hasattr(peak_date, 'strftime') else str(peak_date))} un pico de {int(peak_count):,} consultas. "
            f"Este volumen representa un incremento porcentual (IP%) de {ip_pct:,.1f} % sobre su promedio diario registrado de {vnp:.1f} consultas. "
            f"Este comportamiento representa un pico abrupto y aislado típico de procesos de extracción automatizada (ver imagen {image_num})."
        )
    elif chart_id == "G2":
        hourly = df.copy()
        hourly["hora"] = hourly["fecha_consulta"].dt.hour
        total_consultas = len(hourly)

        if total_consultas > 0:
            hourly_counts = hourly.groupby("hora").size()
            peak_hour = hourly_counts.idxmax()
            peak_hour_count = hourly_counts.max()
            peak_hour_pct = (peak_hour_count / total_consultas) * 100

            # Consultas en horario laboral (08:00 - 17:59)
            biz_count = hourly[(hourly["hora"] >= 8) & (hourly["hora"] < 18)].shape[0]
            biz_pct = (biz_count / total_consultas) * 100

            # Consultas en la madrugada (00:00 - 05:59)
            noc_count = hourly[(hourly["hora"] >= 0) & (hourly["hora"] < 6)].shape[0]
            noc_pct = (noc_count / total_consultas) * 100

            off_pct = 100.0 - biz_pct

            if biz_pct >= 70.0:
                return (
                    f"En el análisis de horarios, la actividad se concentra principalmente dentro del horario laboral regular (08:00 a 18:00 h), "
                    f"acumulando el {biz_pct:.1f}% del total de consultas. El pico máximo se registra a las {peak_hour:02d}:00 horas con "
                    f"{peak_hour_count:,} consultas ({peak_hour_pct:.1f}% del volumen del gráfico). Esta concentración en horas de oficina "
                    f"sugiere un uso intensivo mediante procesos masivos o consultas en ráfagas integradas a la operación diaria (ver imagen {image_num})."
                )
            elif off_pct >= 35.0:
                return (
                    f"En el análisis de horarios, se identifica una concentración atípica fuera del horario operativo tradicional. "
                    f"El {off_pct:.1f}% de las consultas se realizó entre las 18:00 y las 07:59 h (incluyendo un {noc_pct:.1f}% en la madrugada). "
                    f"El pico máximo de consultas se registra a las {peak_hour:02d}:00 horas, acumulando {peak_hour_count:,} consultas. "
                    f"Este patrón fuera de horas de oficina sugiere la ejecución de scripts o tareas de fondo programadas (ver imagen {image_num})."
                )
            else:
                return (
                    f"El análisis temporal por hora revela una distribución continua a lo largo de las 24 horas del día. El pico máximo "
                    f"de consultas se registra a las {peak_hour:02d}:00 horas ({peak_hour_count:,} consultas), manteniéndose un consumo "
                    f"sostenido del {noc_pct:.1f}% durante la madrugada. Esto demuestra la presencia de consultas automatizadas desatendidas "
                    f"que operan de forma ininterrumpida (ver imagen {image_num})."
                )
        else:
            return f"No se registró actividad horaria suficiente para el gráfico (ver imagen {image_num})."
    elif chart_id == "G3":
        weekday_df = df.copy()
        weekday_df["dayofweek"] = weekday_df["fecha_consulta"].dt.dayofweek
        total_consultas = len(weekday_df)

        if total_consultas > 0:
            days_map = {0: "lunes", 1: "martes", 2: "miércoles", 3: "jueves", 4: "viernes", 5: "sábado", 6: "domingo"}
            counts_by_day = weekday_df.groupby("dayofweek").size()

            peak_day_idx = counts_by_day.idxmax()
            peak_day_name = days_map.get(peak_day_idx, "día habitual")
            peak_day_count = counts_by_day.max()
            peak_day_pct = (peak_day_count / total_consultas) * 100

            # Lunes (0) a Viernes (4)
            weekday_count = counts_by_day.reindex([0, 1, 2, 3, 4], fill_value=0).sum()
            weekday_pct = (weekday_count / total_consultas) * 100

            # Sábado (5) y Domingo (6)
            weekend_count = counts_by_day.reindex([5, 6], fill_value=0).sum()
            weekend_pct = (weekend_count / total_consultas) * 100
            sunday_count = counts_by_day.get(6, 0)

            if weekday_pct >= 80.0:
                return (
                    f"Al analizar la distribución por día de la semana, el consumo se concentra principalmente en días laborables (lunes a viernes), "
                    f"acumulando el {weekday_pct:.1f}% del volumen total ({weekday_count:,} consultas). El volumen máximo semanal se registró el día "
                    f"{peak_day_name} con {peak_day_count:,} consultas (representando el {peak_day_pct:.1f}% del total semanal). "
                    f"Asimismo, se observa una actividad sostenida en fines de semana de {weekend_count:,} consultas ({sunday_count:,} el domingo), "
                    f"evidenciando que el uso principal ocurre en días hábiles con alta carga en días centrales de la semana (ver imagen {image_num})."
                )
            else:
                return (
                    f"El análisis semanal revela una operación continua extendida hacia los fines de semana. Si bien el pico más alto ocurrió el día "
                    f"{peak_day_name} con {peak_day_count:,} consultas ({peak_day_pct:.1f}% de la semana), la entidad mantuvo un flujo de "
                    f"{weekend_count:,} consultas los días sábado y domingo ({weekend_pct:.1f}% del total semanal, acumulando {sunday_count:,} consultas los domingos). "
                    f"Esta continuidad operativa evidencia la ejecución de servicios o scripts de consulta desatendidos de fondo (ver imagen {image_num})."
                )
        else:
            return f"No se registró actividad diaria suficiente para el gráfico (ver imagen {image_num})."
    elif chart_id == "G4":
        user_counts = df["cod_usuario"].value_counts()
        if not user_counts.empty:
            top_user_cod = user_counts.index[0]
            top_user_cnt = user_counts.values[0]
            top_user_pct = (top_user_cnt / len(df)) * 100
            top_user_name = df[df["cod_usuario"] == top_user_cod]["nombre_usuario"].iloc[0] if "nombre_usuario" in df.columns else top_user_cod
            second_user_str = ""
            if len(user_counts) > 1:
                u2_cod = user_counts.index[1]
                u2_cnt = user_counts.values[1]
                u2_name = df[df["cod_usuario"] == u2_cod]["nombre_usuario"].iloc[0] if "nombre_usuario" in df.columns else u2_cod
                second_user_str = f", seguido por {u2_name} (DNI: {u2_cod}) con {u2_cnt:,} consultas"
        else:
            top_user_cod, top_user_cnt, top_user_pct, top_user_name, second_user_str = "N/A", 0, 0, "N/A", ""
        return (
            f"La actividad global de la entidad muestra una distribución donde dos cuentas concentran la mayor carga de uso: "
            f"el usuario {top_user_name} (DNI: {top_user_cod}) ejecutó {top_user_cnt:,} consultas ({top_user_pct:.2f}% del total){second_user_str}. "
            f"Esta concentración exige una auditoría directa a los perfiles y atribuciones de negocio de ambos operadores. (ver imagen {image_num})."
        )
    elif chart_id == "G5":
        if not report.bursts.empty:
            max_burst_idx = report.bursts["count"].idxmax()
            max_burst = report.bursts.loc[max_burst_idx]
            max_burst_cnt = max_burst["count"]
            max_burst_time = max_burst["window_start"].strftime("%d/%m/%Y a las %H:%M:%S") if hasattr(max_burst["window_start"], 'strftime') else str(max_burst["window_start"])
            max_burst_user = max_burst["cod_usuario"]
        else:
            max_burst_cnt, max_burst_time, max_burst_user = 0, "N/A", "N/A"
        return (
            f"Se detectaron un total de {len(report.bursts):,} eventos de ráfagas masivas en ventanas de 60 segundos. "
            f"La ráfaga más intensa alcanzó un pico de {max_burst_cnt} consultas en un minuto el día {max_burst_time} por el "
            f"usuario {max_burst_user}, lo que representa una velocidad inaccesible para un operador manual y "
            f"confirma una ejecución automatizada por software (ver imagen {image_num})."
        )
    elif chart_id == "T_BURST":
        if report.bursts is None or report.bursts.empty:
            return "No se registraron ráfagas de velocidad inusual superiores a los umbrales de alerta en ventanas de 60 segundos."

        num_bursts = len(report.bursts)
        top_burst = report.bursts.sort_values(by="count", ascending=False).iloc[0]
        max_cnt = int(top_burst["count"])
        max_speed = max_cnt / 60.0
        top_user = top_burst["cod_usuario"]
        dt_val = top_burst["window_start"]
        dt_str = dt_val.strftime("%d/%m/%Y a las %H:%M:%S") if hasattr(dt_val, "strftime") else str(dt_val)

        if max_speed >= 3.0:
            return (
                f"En el análisis detallado de velocidad de respuesta, se identificaron {num_bursts:,} ráfagas inusuales en ventanas de 60 segundos. "
                f"La ráfaga de mayor intensidad alcanzó un consumo masivo de {max_cnt:,} consultas en solo un minuto ({max_speed:.2f} consultas/segundo), "
                f"registrada el día {dt_str} por el usuario con DNI {top_user}. Esta velocidad de procesamiento supera la capacidad de interacción "
                f"humana regular, confirmando técnicamente el uso de herramientas o scripts automatizados de consulta (ver tabla a continuación)."
            )
        else:
            return (
                f"El desglose de eventos por ráfaga revela un total de {num_bursts:,} intervalos con concentración inusual de consultas en ventanas de 60 segundos. "
                f"El pico de mayor intensidad registró {max_cnt:,} consultas ({max_speed:.2f} consultas/segundo) el día {dt_str} por el usuario "
                f"con DNI {top_user}. Esta concentración de solicitudes consecutivas en periodos de un minuto refleja la automatización de procesos "
                f"de consulta en lote (ver tabla a continuación)."
            )
    elif chart_id == "G6":
        df_hv = df.copy()
        df_hv["hora_periodo"] = df_hv["fecha_consulta"].dt.floor("h")
        hourly_counts = df_hv.groupby("hora_periodo").size()
        if not hourly_counts.empty:
            max_time = hourly_counts.idxmax()
            max_cnt = hourly_counts.max()
            total_cnt = len(df)
            pct_period = (max_cnt / total_cnt * 100) if total_cnt > 0 else 0.0
            cpm = max_cnt / 60.0
            cps = cpm / 60.0
            
            dias_semana_es = {0: "Lunes", 1: "Martes", 2: "Miércoles", 3: "Jueves", 4: "Viernes", 5: "Sábado", 6: "Domingo"}
            dia_nombre = dias_semana_es.get(max_time.weekday(), "") if hasattr(max_time, "weekday") else ""
            
            meses_nombres = {1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio", 
                             7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"}
            day_num = max_time.day if hasattr(max_time, "day") else ""
            month_name = meses_nombres.get(max_time.month, "") if hasattr(max_time, "month") else ""
            year_num = max_time.year if hasattr(max_time, "year") else ""
            hora_fmt = max_time.strftime("%H:00 h") if hasattr(max_time, "strftime") else ""
            
            peak_str = f"{dia_nombre} {day_num} de {month_name} de {year_num} a las {hora_fmt}".strip()
        else:
            max_cnt, pct_period, cpm, cps, peak_str = 0, 0.0, 0.0, 0.0, "N/A"
            
        num_hv = len(report.high_volume) if not report.high_volume.empty else 0
        return (
            f"El análisis de dispersión temporal de volumen inusual registró {num_hv:,} intervalo(s) horario(s) que superaron "
            f"el umbral auditado de 100 consultas/hora. El volumen máximo concentrado en un solo intervalo se registró el "
            f"{peak_str}, alcanzando un pico de {max_cnt:,} consultas en esa sola hora. Esta concentración representó el "
            f"{pct_period:.2f}% de todas las consultas del periodo evaluado, lo que equivale a una velocidad promedio sostenida "
            f"de {cpm:.2f} consultas por minuto (~{cps:.2f} consultas/segundo) durante los 60 minutos ininterrumpidos. "
            f"Esta tasa de transferencia continua confirma una modalidad de extracción masiva desatendida mediante software "
            f"o scripts clientes (ver imagen {image_num})."
        )
    elif chart_id == "G7":
        num_rep = len(report.repeated_dni) if not report.repeated_dni.empty else 0
        if num_rep == 0 and "dni_consultado" in df.columns:
            dni_vc = df["dni_consultado"].value_counts()
            num_rep = len(dni_vc[dni_vc > 1])
        return (
            f"El histograma de frecuencia de DNIs repetidos evidencia la existencia de {num_rep:,} documentos de identidad "
            f"que registran múltiples consultas durante el periodo analizado, superando el umbral de repetición establecido. "
            f"La distribución muestra una concentración atípica en la frecuencia de reconsultas sobre personas específicas, "
            f"lo que sugiere patrones de monitoreo recurrente o consultas persistentes desatendidas (ver imagen {image_num})."
        )
    elif chart_id == "G8":
        dni_counts = df["dni_consultado"].value_counts() if "dni_consultado" in df.columns else pd.Series()
        if not dni_counts.empty:
            max_rep_dni = dni_counts.index[0]
            max_rep_cnt = dni_counts.values[0]
            top_user_info = ""
            if "cod_usuario" in df.columns:
                dni_df = df[df["dni_consultado"] == max_rep_dni]
                if not dni_df.empty:
                    top_u = dni_df["cod_usuario"].value_counts().index[0]
                    top_u_name = dni_df[dni_df["cod_usuario"] == top_u]["nombre_usuario"].iloc[0] if "nombre_usuario" in dni_df.columns else top_u
                    top_user_info = f", registrando mayor actividad por parte del usuario {top_u_name} (DNI: {top_u})"
        else:
            max_rep_dni, max_rep_cnt, top_user_info = "N/A", 0, ""
        return (
            f"El ranking de los DNIs más consultados por los usuarios del servicio evidencia la focalización anómala sobre ciudadanos específicos. "
            f"El documento más consultado corresponde al DNI {max_rep_dni} con un total de {max_rep_cnt:,} consultas{top_user_info}. "
            f"Esta consulta intensiva y reiterada sobre un conjunto acotado de DNIs requiere la verificación inmediata de las justificaciones de negocio de las operaciones realizadas (ver imagen {image_num})."
        )
    elif chart_id == "G9":
        num_periodic = len(report.periodic) if not report.periodic.empty else 0
        return (
            f"El análisis de regularidad temporal detectó {num_periodic} usuario(s) con comportamiento de consulta estrictamente periódico. "
            f"Una desviación estándar en los intervalos de tiempo menor a 2.0 segundos es un indicador inequívoco "
            f"de procesos de consulta automatizados (scripts/bots) programados para extraer datos de forma sistemática y desatendida (ver imagen {image_num})."
        )
    elif chart_id == "G10":
        df_hm = df.copy()
        total_consultas = len(df_hm)
        if total_consultas > 0:
            df_hm["hora"] = df_hm["fecha_consulta"].dt.hour
            df_hm["dayofweek"] = df_hm["fecha_consulta"].dt.dayofweek

            days_map = {0: "lunes", 1: "martes", 2: "miércoles", 3: "jueves", 4: "viernes", 5: "sábado", 6: "domingo"}

            # Agrupar por (día de la semana, hora)
            cell_counts = df_hm.groupby(["dayofweek", "hora"]).size()
            peak_cell_idx = cell_counts.idxmax()
            peak_day_num, peak_hour = peak_cell_idx
            peak_day_name = days_map.get(peak_day_num, "día de la semana")
            peak_cell_count = cell_counts.max()
            peak_cell_pct = (peak_cell_count / total_consultas) * 100

            # Bloques laborales: Lunes(0) a Viernes(4) entre 08:00 y 17:59 h
            biz_mask = (df_hm["dayofweek"] < 5) & (df_hm["hora"] >= 8) & (df_hm["hora"] < 18)
            biz_count = df_hm[biz_mask].shape[0]
            biz_pct = (biz_count / total_consultas) * 100
            off_pct = 100.0 - biz_pct

            if biz_pct >= 65.0:
                return (
                    f"El mapa de calor de intensidad de consultas (Hora × Día) revela una concentración crítica focalizada en la jornada laboral de días hábiles. "
                    f"El bloque de mayor intensidad operativa se registró los días {peak_day_name} a las {peak_hour:02d}:00 horas, alcanzando un pico de "
                    f"{peak_cell_count:,} consultas en esa casilla temporal ({peak_cell_pct:.1f}% del volumen total). El {biz_pct:.1f}% de la densidad calórica "
                    f"coincide con el horario de oficina (08:00 a 18:00 h), lo que demuestra un patrón de uso intensivo o ráfagas masivas integradas a la operación diaria (ver imagen {image_num})."
                )
            elif off_pct >= 35.0:
                return (
                    f"El mapa de calor muestra puntos calientes de alta densidad fuera de la jornada de oficina tradicional. La casilla con el mayor consumo concentrado "
                    f"ocurrió los días {peak_day_name} a las {peak_hour:02d}:00 horas con {peak_cell_count:,} consultas. El {off_pct:.1f}% del consumo se distribuye en "
                    f"bloques nocturnos y/o fines de semana, lo que confirma que la actividad está guiada por scripts desatendidos programados fuera del horario laboral (ver imagen {image_num})."
                )
            else:
                return (
                    f"El mapa de calor ilustra una distribución calórica homogénea y continua que abarca todas las franjas horarias y días de la semana. "
                    f"El bloque con el mayor volumen relativo corresponde a los días {peak_day_name} a las {peak_hour:02d}:00 horas con {peak_cell_count:,} consultas. "
                    f"Este patrón de manchas térmicas constantes durante las 24 horas evidencia la operación ininterrumpida de servicios de consulta desatendidos (ver imagen {image_num})."
                )
        else:
            return f"No se registró actividad temporal suficiente para generar el mapa de calor (ver imagen {image_num})."
    return f"Visualización analítica de los consumos (ver imagen {image_num})."

def build_report(
    df,
    metadata: dict,
    report,
    charts: dict[str, str],
    dni_infos: list,
    template_path: str,
    output_path: str
) -> str:
    """
    Copia la plantilla Word y reemplaza todos los campos con los resultados de las anomalías y los gráficos.
    """
    logger.info("Iniciando construcción del informe Word con nuevo formato...")
    out_dir = Path(output_path)
    out_dir.mkdir(parents=True, exist_ok=True)
    
    empresa_clean = re.sub(r'[\\/*?:"<>|]', '', str(metadata.get("empresa_principal", "ENTIDAD"))).strip()
    
    meses_upper = {1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL", 5: "MAYO", 6: "JUNIO", 
                   7: "JULIO", 8: "AGOSTO", 9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE"}
    end_date = metadata.get("fecha_fin")
    if isinstance(end_date, datetime) or hasattr(end_date, "month"):
        mes_str = meses_upper.get(end_date.month, "MES")
    else:
        mes_str = "MES"
        
    serv_raw = str(metadata.get("servicio", "")).upper()
    if "LINEA DEDICADA" in serv_raw or "LD" in serv_raw:
        serv_abbrev = "LD"
    elif "INTERNET" in serv_raw:
        serv_abbrev = "INTERNET"
    elif "FACIAL" in serv_raw:
        serv_abbrev = "WSBF"
    elif "DACTILAR" in serv_raw:
        serv_abbrev = "WSBD"
    elif "BIOMETRI" in serv_raw:
        serv_abbrev = "BIOMETRIA"
    elif "WEB SERVICE" in serv_raw or "WS" in serv_raw:
        serv_abbrev = "WS"
    else:
        serv_abbrev = "SERV"

    doc_base_name = f"{empresa_clean}-{mes_str}-{serv_abbrev}"
    final_docx_path = out_dir / f"{doc_base_name}.docx"
    
    if final_docx_path.exists():
        time_suffix = datetime.now().strftime("%H%M%S")
        final_docx_path = out_dir / f"{doc_base_name}_{time_suffix}.docx"

    # 1. El gráfico G1 (Imagen 1) se genera dinámicamente en chart_generator a partir de los SADAM de data/
    #    (ya no se reutiliza el PNG histórico de REVISION SEMANAL DE CONSUMOS, que queda desactualizado).

    # 2. Obtener métricas históricas de consumo (VNP 2025-2026, VID, IP%, Peak Date)
    hist_metrics = get_historical_metrics(metadata["empresa_principal"], metadata["ruc_principal"], metadata["servicio"], df)
    vnp = hist_metrics["vnp"]
    vid = hist_metrics["vid"]
    ip_pct = hist_metrics["ip_pct"]
    peak_date = hist_metrics["peak_date"]

    df_copy = df.copy()
    df_copy["fecha_dia"] = df_copy["fecha_consulta"].dt.date
    daily_user = df_copy.groupby(["cod_usuario", "fecha_dia"]).size().reset_index(name="count")
    
    if not daily_user.empty:
        idx_max = daily_user["count"].idxmax()
        peak_row = daily_user.loc[idx_max]
        peak_user_cod = peak_row["cod_usuario"]
        peak_user_name = peak_user_cod
        if "nombre_usuario" in df.columns:
            names = df_copy[df_copy["cod_usuario"] == peak_user_cod]["nombre_usuario"].dropna().unique()
            if len(names) > 0:
                peak_user_name = names[0]
        peak_user_full = f"{peak_user_name} (DNI: {peak_user_cod})"
    else:
        peak_user_cod = "SIN_USUARIO"
        peak_user_name = "SIN_USUARIO"
        peak_user_full = "SIN_USUARIO"

    # Formatear lista de usuarios principales para la cabecera y selección del usuario que más consume
    user_counts_all = df_copy["cod_usuario"].value_counts()
    if not user_counts_all.empty:
        top_user_cod = user_counts_all.index[0]
        top_user_name = top_user_cod
        if "nombre_usuario" in df.columns:
            names = df_copy[df_copy["cod_usuario"] == top_user_cod]["nombre_usuario"].dropna().unique()
            if len(names) > 0:
                top_user_name = names[0]
        top_user_eval = f"{top_user_name} (DNI: {top_user_cod})"
    else:
        top_user_eval = "SIN_USUARIO"

    formatted_anom_users = []
    for u_cod in user_counts_all.index[:2]:
        u_df = df_copy[df_copy["cod_usuario"] == u_cod]
        u_name = u_cod
        if "nombre_usuario" in df.columns:
            names = u_df["nombre_usuario"].dropna().unique()
            if len(names) > 0:
                u_name = names[0]
        formatted_anom_users.append(f"{u_name} (DNI: {u_cod})")
    usuarios_str = " / ".join(formatted_anom_users)

    # 1.1 Obtener cantidad total de usuarios del servicio desde USUARIOS POR ENTIDAD.xlsx
    service_user_count = get_service_user_count(
        empresa_name=metadata["empresa_principal"],
        service_name=metadata["servicio"],
        ruc=metadata.get("ruc_principal")
    )
    if service_user_count is not None:
        user_count_val = service_user_count
    else:
        user_count_val = df_copy["cod_usuario"].nunique() if not df_copy.empty else 0

    # Determinar el periodo evaluado
    meses = {1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio", 
             7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre", 11: "noviembre", 12: "diciembre"}
    end_dt = metadata.get("fecha_fin")
    current_month_num = end_dt.month if isinstance(end_dt, datetime) and hasattr(end_dt, "month") else datetime.now().month
    current_month_name = meses.get(current_month_num, "julio").capitalize()
    periodo_evaluado_str = f"Enero - {current_month_name} 2026"

    # Copiar plantilla Word
    t_path = Path(template_path)
    if not t_path.exists():
        raise FileNotFoundError(f"No se encontró la plantilla en {template_path}")
        
    shutil.copy(template_path, final_docx_path)
    doc = docx.Document(final_docx_path)

    # 3. Generar las 3 Conclusiones Justificadas y Extensas
    num_bursts = len(report.bursts) if not report.bursts.empty else 0
    max_burst_cnt = report.bursts["count"].max() if not report.bursts.empty else 0
    max_burst_speed = (max_burst_cnt / 60.0) if max_burst_cnt > 0 else 0.0

    top_user_volume = user_counts_all.iloc[0] if not user_counts_all.empty else 0
    total_filas_df = len(df)
    top_user_pct = (top_user_volume / total_filas_df * 100) if total_filas_df > 0 else 0.0

    max_rep_dni = "N/A"
    max_rep_cnt = 0
    if not report.repeated_dni.empty:
        max_rep_dni = report.repeated_dni.iloc[0]["dni_consultado"]
        max_rep_cnt = report.repeated_dni.iloc[0]["count"]
    elif "dni_consultado" in df.columns:
        dni_counts = df["dni_consultado"].value_counts()
        if not dni_counts.empty:
            max_rep_dni = dni_counts.index[0]
            max_rep_cnt = dni_counts.values[0]

    if num_bursts > 0:
        concl_1 = (
            f"Evidencia de Automatización: La detección de {num_bursts:,} ráfagas masivas y un pico máximo de "
            f"{max_burst_cnt} consultas en 60 segundos ({max_burst_speed:.1f} consultas/segundo) presume técnicamente el uso de "
            f"software de consulta automatizado o scripts sin restricción de tasa de peticiones."
        )
    else:
        num_illegal = len(report.illegal_hours) if not report.illegal_hours.empty else 0
        concl_1 = (
            f"Evidencia de Automatización: El análisis temporal e intervalos en la actividad consumida por la entidad "
            f"revela la ejecución desatendida de procesos automatizados, registrándose {num_illegal:,} consultas en franjas horarias "
            f"no operativas (madrugada y/o fines de semana), lo que sugiere el uso de scripts o bots desvinculados de la atención presencial regular."
        )

    num_accts = len(formatted_anom_users)
    if num_accts == 1:
        acct_str_phrase = f"por la cuenta de usuario ({usuarios_str})"
    else:
        acct_str_phrase = f"por {num_accts} cuentas de usuario ({usuarios_str})"

    concl_2 = (
        f"Concentración Extrema del Consumo: Más de {top_user_volume:,} consultas (más del {top_user_pct:.1f}% del total "
        f"institucional de {total_filas_df:,} filas) fueron ejecutadas {acct_str_phrase}, "
        f"situando en ella(s) la prioridad del control interno."
    )

    concl_3 = (
        f"Reincidencia Anómala en Documentos de Identidad: La consulta de un mismo DNI hasta {max_rep_cnt:,} veces en el mes "
        f"(DNI {max_rep_dni}) sugiere la filtración, monitoreo continuo o fallas en el software de integración cliente de "
        f"{metadata['empresa_principal']}, violando la finalidad del servicio."
    )

    conclusiones_list = [concl_1, concl_2, concl_3]

    # 4. Generar la Recomendación Justificada y Detallada
    recomendacion_text = (
        f"Se recomienda evaluar el bloqueo preventivo inmediato del código de usuario y credenciales asociadas a {usuarios_str}, "
        f"en los sistemas de autenticación de RENIEC. Paralelamente, se formalizará una exigencia de descargos a "
        f"{metadata['empresa_principal']} para que su Oficial de Seguridad de la Información remita en un plazo no mayor "
        f"a cinco (5) días hábiles el informe interno detallando la justificación de negocio, la IP origen de la "
        f"petición interna y la naturaleza del proceso asociado a dicho usuario."
    )

    # 5. Reemplazar placeholders en el cuerpo
    fecha_hoy_str = datetime.now().strftime("%d-%m-%Y")
    replace_placeholder(doc, "{{EMPRESA}}", metadata["empresa_principal"])
    replace_placeholder(doc, "{{RUC}}", metadata["ruc_principal"])
    replace_placeholder(doc, "{{SERVICIO}}", metadata["servicio"])
    replace_placeholder(doc, "{{TOTAL_CONSULTAS}}", f"{metadata['total_filas']:,}")

    # Actualizar campos de Página 1 (Información General)
    for p in doc.paragraphs[:20]:
        if "Usuarios del Servicio:" in p.text or "Cantidad de usuarios del Servicio:" in p.text:
            p.text = ""
            p.add_run(f"Cantidad de usuarios del Servicio: {user_count_val} usuarios")
            if user_count_val == 1:
                r_gen = p.add_run(" (Usuario genérico)")
                r_gen.bold = True
                r_gen.font.color.rgb = RGBColor(237, 28, 36)

    mes_evaluado_lower = current_month_name.lower()
    replace_placeholder(doc, "en el mes de mayo", f"en el mes de {mes_evaluado_lower}")
    replace_placeholder(doc, "el mes de mayo", f"el mes de {mes_evaluado_lower}")
    replace_placeholder(doc, "CRISTHIAN ALFONSO LUJAN MENDOZA (DNI: 41002681)", top_user_eval)
    replace_placeholder(doc, "28-05-2026", fecha_hoy_str)
    replace_placeholder(doc, "29-05-2026", fecha_hoy_str)
    replace_placeholder(doc, "enero - mayo 2026", periodo_evaluado_str)
    replace_placeholder(doc, "enero – mayo 2026", periodo_evaluado_str)
    replace_placeholder(doc, "Enero - Julio 2026", periodo_evaluado_str)

    # 6. Reestructurar la Sección 4.2 y subsecciones secuencialmente
    possible_charts = [
        ("G1", f"Análisis del Consumo por Usuario (Enero – {current_month_name} 2026)"),
        ("G2", "Distribución de Consultas por Hora del Día"),
        ("G3", "Actividad de Consultas por Día de la Semana"),
        ("G4", "Distribución por Usuario/Código"),
        ("G5", "Frecuencia de Ráfagas por Hora del Día"),
        ("T_BURST", "Top 10 Ráfagas de Mayor Intensidad en 60 Segundos"),
        ("G6", "Dispersión de Volumen Horario Inusual"),
        ("G8", "Top 10 DNIs Más Consultados por los Usuarios del Servicio"),
        ("G9", "Análisis de Intervalos entre Consultas"),
        ("G10", "Mapa de Calor de Consumo (Hora × Día)")
    ]

    # Encontrar "4.2.1" o "(Imagen 1)" y limpiar los párrafos posteriores de la plantilla antigua
    idx_img1 = -1
    for i, p in enumerate(doc.paragraphs):
        if "4.2.1" in p.text or "(Imagen 1)" in p.text:
            idx_img1 = i
            break

    if idx_img1 != -1:
        elems_to_del = [doc.paragraphs[k]._element for k in range(idx_img1, len(doc.paragraphs))]
        for el in elems_to_del:
            el.getparent().remove(el)

    item_counter = 1
    img_counter = 1
    sub_counter = 1
    for chart_id, title_raw in possible_charts:
        if chart_id == "T_BURST":
            if report.bursts is None or report.bursts.empty:
                continue
        elif chart_id == "G6":
            if user_count_val != 1:
                continue
            if chart_id not in charts or not os.path.exists(charts[chart_id]):
                continue
        elif chart_id not in charts or not os.path.exists(charts[chart_id]):
            continue

        title_str = f"4.2.{sub_counter} {title_raw}"
        sub_counter += 1

        # Salto de página a partir del segundo elemento para garantizar 1 página por punto
        if item_counter > 1:
            doc.add_page_break()

        # Título de subsección alineado a la izquierda sin sangría
        p_title = doc.add_paragraph()
        format_heading_left(p_title)
        run_title = p_title.add_run(f"\n{title_str}\n" if item_counter == 1 else f"{title_str}\n")
        run_title.bold = True
        p_title.paragraph_format.space_before = Pt(6)
        p_title.paragraph_format.space_after = Pt(4)

        # Si es G1, agregamos la etiqueta de Usuario antes del párrafo explicativo
        if chart_id == "G1":
            p_user_label = doc.add_paragraph()
            format_heading_left(p_user_label)
            p_user_label.add_run(f"Usuario: {peak_user_full}\n")

        # Párrafo explicativo / hallazgo justificado
        p_hall = doc.add_paragraph()
        p_hall.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        hallazgo_txt = obtener_hallazgo_dinamico(chart_id, df, metadata, report, peak_date, vid, vnp, ip_pct, peak_user_name, peak_user_cod, image_num=img_counter)
        p_hall.add_run(hallazgo_txt)
        p_hall.paragraph_format.space_after = Pt(6)

        if chart_id == "T_BURST":
            add_top_bursts_table(doc, df, report, top_n=10)
        else:
            # Imagen centrada
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            # Ancho óptimo de 14.5 cm (5.71 pulgadas) para ajustar en 1 página
            run_img.add_picture(charts[chart_id], width=Inches(14.5 / 2.54))

            # Etiqueta (Imagen X) centrada debajo del gráfico
            p_label = doc.add_paragraph()
            p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_label.add_run(f"(Imagen {img_counter})")
            p_label.paragraph_format.space_after = Pt(12)
            img_counter += 1

        item_counter += 1

    # 7. Agregar 4.3 Conclusiones (3 conclusiones justificadas)
    doc.add_paragraph()
    p_concl_head = doc.add_paragraph()
    format_heading_left(p_concl_head)
    run_ch = p_concl_head.add_run("4.3. Conclusiones:\n")
    run_ch.bold = True

    for c_text in conclusiones_list:
        p_c = doc.add_paragraph()
        p_c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_c.add_run(c_text)
        p_c.paragraph_format.space_after = Pt(6)

    # 8. Agregar 4.4 Recomendaciones (1 recomendación justificada)
    doc.add_paragraph()
    p_recom_head = doc.add_paragraph()
    format_heading_left(p_recom_head)
    run_rh = p_recom_head.add_run("4.4. Recomendaciones:\n")
    run_rh.bold = True

    p_r = doc.add_paragraph()
    p_r.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_r.add_run(recomendacion_text)
    p_r.paragraph_format.space_after = Pt(6)

    # 9. Actualizar Tabla 0 (Resumen en primera hoja)
    if len(doc.tables) > 0:
        table = doc.tables[0]
        if len(table.rows) >= 4 and len(table.columns) >= 4:
            logger.info("Actualizando Tabla de Resumen de Consumos (Tabla 0)...")
            table.cell(1, 0).text = str(metadata["servicio"])
            table.cell(1, 1).text = str(usuarios_str)
            table.cell(1, 2).text = "Volumen diario Promedio (VNP)"
            table.cell(1, 3).text = f"{vnp:.1f}"
            
            table.cell(2, 0).text = str(metadata["servicio"])
            table.cell(2, 1).text = str(usuarios_str)
            table.cell(2, 2).text = "Volumen Inusual Detectado (VID)"
            table.cell(2, 3).text = f"{int(vid):,}"
            
            table.cell(3, 0).text = str(metadata["servicio"])
            table.cell(3, 1).text = str(usuarios_str)
            table.cell(3, 2).text = "Incremento Porcentual (IP %)"
            table.cell(3, 3).text = f"{ip_pct:,.1f} %"
            
            for row_idx in range(1, 4):
                set_cell_background(table.cell(row_idx, 2), "F2F2F2")

    doc.save(final_docx_path)
    logger.info(f"Informe Word generado exitosamente en: {final_docx_path}")
    return str(final_docx_path.resolve())
